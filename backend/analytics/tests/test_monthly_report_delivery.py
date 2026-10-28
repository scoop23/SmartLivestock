"""Exercise report generation and delivery with isolated data and local email."""

from datetime import date
from io import BytesIO, StringIO
from unittest.mock import patch

from django.core import mail
from django.core.management import call_command, CommandError
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from analytics.models import MonthlyReportDelivery
from users.models import Role, User


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class MonthlyReportDeliveryTests(TestCase):
    def setUp(self):
        role = Role.objects.create(role_name="MAO")
        User.objects.create_user(
            username="audit-mao", email="audit-mao@example.invalid", password=None,
            role=role, account_status=User.AccountStatus.APPROVED,
        )

    def send(self):
        with patch.dict("os.environ", {"AUTOMATED_REPORT_RECIPIENTS": ""}):
            call_command("send_monthly_report", month="2026-09", stdout=StringIO())

    def test_first_delivery_generates_real_workbook_and_duplicate_run_does_not_resend(self):
        self.send()
        delivery = MonthlyReportDelivery.objects.get(period_start=date(2026, 9, 1))
        self.assertEqual(delivery.status, "SENT")
        self.assertIsNotNone(delivery.sent_at)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["audit-mao@example.invalid"])
        workbook = load_workbook(BytesIO(mail.outbox[0].attachments[0].content))
        self.assertIn("production", workbook.sheetnames)
        self.send()
        self.assertEqual(len(mail.outbox), 1)

    def test_existing_running_delivery_is_not_taken_over(self):
        MonthlyReportDelivery.objects.create(period_start=date(2026, 9, 1), period_end=date(2026, 9, 30))
        with self.assertRaisesMessage(CommandError, "already being generated"):
            self.send()

    def test_email_failure_is_recorded_and_retry_can_complete(self):
        with patch("analytics.management.commands.send_monthly_report.EmailMessage.send", side_effect=OSError("synthetic outage")):
            with self.assertRaises(CommandError):
                self.send()
        self.assertEqual(MonthlyReportDelivery.objects.get().status, "FAILED")
        self.send()
        self.assertEqual(MonthlyReportDelivery.objects.get().status, "SENT")
