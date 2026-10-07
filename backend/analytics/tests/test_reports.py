from datetime import time, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient
from openpyxl import load_workbook

from diseases.models import DiseaseCase, MortalityRecord
from livestock.models import Barangay, Farmer, LivestockInventory, LivestockType
from movements.models import LivestockInspection, LivestockInspectionClearance, LivestockInspectionItem
from production.models import LiveAnimalSale, ProductionRecord, SlaughterRecord
from analytics.services.report_exports import _monthly_axis_ticks, _report_chart_flowable, _report_chart_specs
from analytics.services.reports import _build_analysis
from users.models import Role, User


class OfficialReportApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.mao_role, _ = Role.objects.get_or_create(role_name="MAO")
        self.farmer_role, _ = Role.objects.get_or_create(role_name="FARMER")
        self.mao = User.objects.create_user(username="report_mao", email="report_mao@example.test",
                                            password="pass", role=self.mao_role,
                                            account_status=User.AccountStatus.APPROVED)
        self.farmer_user = User.objects.create_user(username="report_farmer", email="report_farmer@example.test",
                                                    password="pass", role=self.farmer_role,
                                                    first_name="Juan", last_name="Dela Cruz",
                                                    account_status=User.AccountStatus.APPROVED)
        self.barangay = Barangay.objects.create(barangay_name="Poblacion", latitude=Decimal("13.88"),
                                               longitude=Decimal("121.21"))
        self.farmer = Farmer.objects.create(user=self.farmer_user, barangay=self.barangay,
                                            address="Purok 1, Poblacion", farm_size=Decimal("1.00"))
        self.cattle = LivestockType.objects.create(name="Cattle")
        self.goat = LivestockType.objects.create(name="Goat")
        self.animal = LivestockInventory.objects.create(
            farmer=self.farmer, livestock_type=self.cattle, tag_number="COW-001", quantity=1,
            status=LivestockInventory.StatusType.APPROVED,
            operational_status=LivestockInventory.OperationalStatus.ACTIVE, created_by=self.farmer_user,
        )
        self.today = timezone.localdate()
        self.client.force_authenticate(self.mao)

        ProductionRecord.objects.create(livestock=self.animal, production_type="MILK", quantity=Decimal("12.5"),
                                        unit="LITERS", record_date=self.today, status="APPROVED", created_by=self.farmer_user)
        ProductionRecord.objects.create(livestock=self.animal, production_type="EGGS", quantity=Decimal("8"),
                                        unit="PIECES", record_date=self.today, status="APPROVED", created_by=self.farmer_user)
        ProductionRecord.objects.create(livestock=self.animal, production_type="MILK", quantity=Decimal("99"),
                                        unit="LITERS", record_date=self.today, status="PENDING", created_by=self.farmer_user)
        DiseaseCase.objects.create(livestock=self.animal, name="Foot and mouth disease", affected_count=2,
                                   record_date=self.today, status="APPROVED", created_by=self.farmer_user)
        DiseaseCase.objects.create(livestock=self.animal, name="Unreviewed", affected_count=20,
                                   record_date=self.today, status="SUBJECT_TO_REVISION", created_by=self.farmer_user)
        MortalityRecord.objects.create(livestock=self.animal, cause="Other", death_count=1,
                                       record_date=self.today, status="APPROVED", created_by=self.farmer_user)
        SlaughterRecord.objects.create(livestock=self.animal, livestock_type=self.cattle, barangay=self.barangay,
                                       quantity=1, carcass_weight=Decimal("180.5"), record_date=self.today,
                                       status="APPROVED", created_by=self.farmer_user)
        SlaughterRecord.objects.create(livestock=self.animal, livestock_type=self.cattle, quantity=9,
                                       carcass_weight=Decimal("999"), record_date=self.today,
                                       status="PENDING", created_by=self.farmer_user)
        LiveAnimalSale.objects.create(livestock=self.animal, quantity=4, sale_date=self.today,
                                      destination="Padre Garcia Livestock Auction Market",
                                      purpose="BREEDING", status="APPROVED", created_by=self.mao)

        self.inspection = LivestockInspection.objects.create(
            shipper=self.farmer, shipper_name="Juan Dela Cruz", destination="Padre Garcia Auction",
            purpose="SLAUGHTER", inspection_date=self.today, created_by=self.mao,
        )
        LivestockInspectionItem.objects.create(inspection=self.inspection, livestock_type=self.cattle,
                                               inventory=self.animal, quantity=1, classification="SLAUGHTER")
        LivestockInspectionItem.objects.create(inspection=self.inspection, livestock_type=self.goat,
                                               quantity=2, classification="OTHER")
        LivestockInspectionClearance.objects.create(
            inspection=self.inspection, control_number="CLR-REPORT-1", shipper_address=self.farmer.address,
            origin="Poblacion, Padre Garcia", status="APPROVED", issued_by=self.mao,
            date_issued=self.today, time_issued=time(9, 0), reviewed_by=self.mao, reviewed_at=timezone.now(),
        )
        pending = LivestockInspection.objects.create(shipper=self.farmer, shipper_name="Juan Dela Cruz",
                                                      destination="Lipa", purpose="SLAUGHTER",
                                                      inspection_date=self.today, created_by=self.mao)
        LivestockInspectionItem.objects.create(inspection=pending, livestock_type=self.cattle,
                                               quantity=10, classification="SLAUGHTER")
        LivestockInspectionClearance.objects.create(inspection=pending, control_number="CLR-REPORT-2",
                                                    shipper_address=self.farmer.address, status="PENDING")

    def params(self, report_type, **overrides):
        values = {"report_type": report_type, "date_from": self.today.isoformat(), "date_to": self.today.isoformat()}
        values.update(overrides)
        return values

    def preview(self, report_type, **filters):
        return self.client.get(reverse("reports-preview"), self.params(report_type, **filters))

    def test_each_report_type_uses_approved_database_rows_and_correct_units(self):
        expected = {"inventory": (1, 1), "production": (2, 2), "disease_mortality": (2, 2),
                    "slaughter": (1, 1), "movement": (3, 2), "inspection": (2, 1)}
        for report_type, (row_count, records) in expected.items():
            with self.subTest(report_type=report_type):
                response = self.preview(report_type)
                self.assertEqual(response.status_code, 200, response.data)
                self.assertEqual(len(response.data["rows"]), row_count)
                self.assertEqual(response.data["record_count"], records)
                self.assertEqual(response.data["period"]["date_from"], self.today.isoformat())
                self.assertIn("executive_summary", response.data["analysis"])
                self.assertIn("key_findings", response.data["analysis"])
                self.assertIn("comparison", response.data["analysis"])
                self.assertIn("coverage_notice", response.data["analysis"])
        production = self.preview("production").data
        self.assertEqual(production["summary"]["totals_by_type_and_unit"], {"MILK (LITERS)": 12.5, "EGGS (PIECES)": 8.0})
        production_kpis = {item["label"]: item["value"] for item in production["analysis"]["executive_summary"]}
        self.assertEqual(production_kpis["Total production (LITERS)"], 12.5)
        self.assertEqual(production_kpis["Total production (PIECES)"], 8)
        self.assertEqual(production_kpis["Contributing farmers"], 1)
        self.assertNotIn("Total production", production_kpis)
        self.assertTrue(production["analysis"]["key_findings"])
        self.assertEqual(production["analysis"]["rankings"][0]["unit"], "LITERS")
        self.assertEqual(production["analysis"]["coverage_level"], "sparse")
        health = self.preview("disease_mortality").data["summary"]
        self.assertEqual((health["disease_cases"], health["affected_heads"], health["mortality_records"], health["deaths"]), (1, 2, 1, 1))
        disease_report = self.preview("disease_mortality").data
        self.assertTrue(any("Foot and mouth disease" in finding for finding in disease_report["analysis"]["key_findings"]))
        movement = self.preview("movement").data
        self.assertEqual(movement["summary"]["livestock_heads"], 3)
        self.assertEqual(movement["summary"]["auction_activity"]["auction_items_processed"], 4)
        self.assertEqual(sum(row["record_kind"] == "AUCTION" for row in movement["rows"]), 1)
        self.assertEqual(sum(row["record_kind"] == "MOVEMENT" for row in movement["rows"]), 2)
        self.assertEqual({row["direction"] for row in movement["rows"] if row["record_kind"] == "MOVEMENT"}, {"INTERNAL"})
        self.assertTrue(all(row["clearance_status"] == "APPROVED" for row in movement["rows"]))

    def test_filters_and_empty_date_range_are_applied_to_preview(self):
        response = self.preview("movement", species="Goat", direction="INTERNAL", purpose="SLAUGHTER")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["rows"]), 1)
        self.assertEqual(response.data["rows"][0]["species"], "Goat")
        self.assertEqual(response.data["rows"][0]["quantity"], 2)
        empty = self.client.get(reverse("reports-preview"), self.params("inventory", date_from="2020-01-01", date_to="2020-01-31"))
        self.assertEqual(empty.status_code, 200)
        self.assertEqual(empty.data["rows"], [])
        self.assertEqual(empty.data["summary"]["total_heads"], 0)

    def test_inventory_registration_dates_return_valid_rows_or_valid_empty_report(self):
        # The model stores a registration timestamp, not snapshots of past holdings.
        historical = self.preview("inventory", date_from="2000-01-01", date_to=self.today.isoformat())
        self.assertEqual(historical.status_code, 200, historical.data)
        self.assertEqual(len(historical.data["rows"]), 1)
        self.assertEqual(historical.data["period_label"], "Registration period")
        self.assertFalse(historical.data["analysis"]["comparison"]["available"])
        self.assertIn("does not store inventory snapshots", historical.data["analysis"]["comparison"]["unavailable_reason"].lower())
        self.assertIn("historical inventory snapshots", historical.data["analysis"]["methodology"].lower())
        self.assertIn("not inferred", historical.data["analysis"]["coverage_notice"].lower())
        inventory_trend = next(chart for chart in _report_chart_specs(historical.data)
                               if chart["title"] == "Historical Inventory Trend")
        self.assertEqual(inventory_trend["type"], "empty")
        self.assertEqual(inventory_trend["data"], [])

        month_start = self.today.replace(day=1)
        current_month = self.preview("inventory", date_from=month_start.isoformat(), date_to=self.today.isoformat())
        self.assertEqual(current_month.status_code, 200, current_month.data)
        self.assertEqual(len(current_month.data["rows"]), 1)

        known_date = self.preview("inventory", date_from=self.today.isoformat(), date_to=self.today.isoformat(),
                                  species="Cattle", barangay="Poblacion")
        self.assertEqual(known_date.status_code, 200, known_date.data)
        self.assertEqual(len(known_date.data["rows"]), 1)

        empty = self.preview("inventory", date_from="2000-01-01", date_to="2000-01-31")
        self.assertEqual(empty.status_code, 200, empty.data)
        self.assertEqual(empty.data["rows"], [])
        self.assertIn("No approved livestock registrations", empty.data["analysis"]["coverage_notice"])

        pending = LivestockInventory.objects.create(
            farmer=self.farmer, livestock_type=self.cattle, tag_number="COW-PENDING", quantity=1,
            status=LivestockInventory.StatusType.PENDING,
            operational_status=LivestockInventory.OperationalStatus.ACTIVE, created_by=self.farmer_user,
        )
        approved_only = self.preview("inventory", date_from="2000-01-01", date_to=self.today.isoformat())
        self.assertEqual(len(approved_only.data["rows"]), 1)
        self.assertNotIn(pending.tag_number, [row["tag_id"] for row in approved_only.data["rows"]])

        missing_dates = self.client.get(reverse("reports-preview"), {"report_type": "inventory"})
        self.assertEqual(missing_dates.status_code, 400)

    def test_report_service_exception_has_safe_server_error_response(self):
        with self.assertLogs("analytics.views", level="ERROR"):
            with patch("analytics.views.build_report", side_effect=RuntimeError("internal diagnostic text")):
                response = self.client.get(reverse("reports-preview"), self.params("inventory"))
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.data, {"detail": "The report service encountered an error."})
        self.assertNotIn("internal diagnostic text", str(response.data))

    def test_date_validation_and_role_authorization(self):
        bad_range = self.client.get(reverse("reports-preview"), self.params("inventory", date_from="2026-02-01", date_to="2026-01-01"))
        self.assertEqual(bad_range.status_code, 400)
        unsupported_filter = self.client.get(reverse("reports-preview"), self.params("inventory", direction="INBOUND"))
        self.assertEqual(unsupported_filter.status_code, 400)
        self.client.force_authenticate(self.farmer_user)
        self.assertEqual(self.preview("inventory").status_code, 403)
        self.client.force_authenticate(user=None)
        self.assertEqual(self.preview("inventory").status_code, 401)

    def test_excel_and_pdf_exports_are_complete_for_every_report(self):
        report_types = ("inventory", "production", "disease_mortality", "slaughter", "movement", "inspection")
        for report_type in report_types:
            with self.subTest(report_type=report_type):
                preview = self.preview(report_type)
                self.assertEqual(preview.status_code, 200, preview.data)
                report = preview.data

                excel = self.client.get(reverse("reports-export"), self.params(report_type, file_format="xlsx"))
                self.assertEqual(excel.status_code, 200)
                self.assertTrue(excel.content.startswith(b"PK"))
                self.assertIn("spreadsheetml.sheet", excel["Content-Type"])
                self.assertIn("attachment; filename=", excel["Content-Disposition"])
                workbook = load_workbook(BytesIO(excel.content), read_only=True)
                sheet = workbook.active
                values = [row for row in sheet.iter_rows(values_only=True)]
                flat_values = {cell for row in values for cell in row if isinstance(cell, str)}
                self.assertIn("EXECUTIVE SUMMARY", flat_values)
                self.assertIn("KEY FINDINGS", flat_values)
                self.assertIn("DATA & METHODOLOGY", flat_values)
                self.assertIn("CHART DATA", flat_values)
                expected_header = tuple(column["label"] for column in report["columns"])
                header_index = next(index for index, row in enumerate(values) if row[:len(expected_header)] == expected_header)
                self.assertEqual(len(values) - header_index - 1, len(report["rows"]))
                for row in report["rows"]:
                    self.assertIn(str(next(iter(row.values()))), flat_values)
                workbook.close()

                pdf = self.client.get(reverse("reports-export"), self.params(report_type, file_format="pdf"))
                self.assertEqual(pdf.status_code, 200)
                self.assertTrue(pdf.content.startswith(b"%PDF"))
                self.assertGreater(len(pdf.content), 500)
                self.assertIn("application/pdf", pdf["Content-Type"])

    def test_empty_excel_export_has_headers_and_empty_state(self):
        params = self.params("inventory", date_from="2000-01-01", date_to="2000-01-31")
        response = self.client.get(reverse("reports-export"), {**params, "file_format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        values = [row for row in workbook.active.iter_rows(values_only=True)]
        self.assertIn("No approved records match this period and filter selection.",
                      {cell for row in values for cell in row if isinstance(cell, str)})
        self.assertIn(tuple(column["label"] for column in self.preview("inventory", date_from="2000-01-01",
                                                                         date_to="2000-01-31").data["columns"]), values)
        workbook.close()
        pdf = self.client.get(reverse("reports-export"), {**params, "file_format": "pdf"})
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_export_filters_match_preview_filters(self):
        filters = {"species": "Goat", "direction": "INTERNAL", "purpose": "SLAUGHTER"}
        preview = self.preview("movement", **filters)
        response = self.client.get(reverse("reports-export"), self.params("movement", file_format="xlsx", **filters))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        columns = tuple(column["label"] for column in preview.data["columns"])
        header_index = next(index for index, row in enumerate(rows) if row[:len(columns)] == columns)
        self.assertEqual(len(rows) - header_index - 1, len(preview.data["rows"]))
        workbook.close()

    def test_export_ignores_preview_page_limits_and_writes_every_matching_row(self):
        ProductionRecord.objects.bulk_create([
            ProductionRecord(livestock=self.animal,
                             production_type="REGRESSION-LAST-ROW" if index == 204 else "MILK",
                             quantity=Decimal("1"),
                             unit="LITERS", record_date=self.today, status="APPROVED", created_by=self.farmer_user)
            for index in range(205)
        ])
        preview = self.preview("production")
        self.assertEqual(len(preview.data["rows"]), 207)
        self.assertEqual(preview.data["rows"][-1]["production_type"], "REGRESSION-LAST-ROW")

        excel = self.client.get(reverse("reports-export"), self.params("production", file_format="xlsx"))
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(BytesIO(excel.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        header = tuple(column["label"] for column in preview.data["columns"])
        header_index = next(index for index, row in enumerate(rows) if row[:len(header)] == header)
        self.assertEqual(len(rows) - header_index - 1, len(preview.data["rows"]))
        workbook.close()

        pdf = self.client.get(reverse("reports-export"), self.params("production", file_format="pdf"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_pdf_chart_specs_use_filtered_rows_and_keep_production_units_separate(self):
        report = self.preview("production", barangay="Poblacion").data
        charts = _report_chart_specs(report)
        titles = [chart["title"] for chart in charts]
        self.assertIn("Production Over Time (LITERS)", titles)
        self.assertIn("Production Over Time (PIECES)", titles)
        milk = next(chart for chart in charts if chart["title"] == "Production by Type (LITERS)")
        eggs = next(chart for chart in charts if chart["title"] == "Production by Type (PIECES)")
        self.assertEqual(milk["data"], [("MILK", 12.5)])
        self.assertEqual(eggs["data"], [("EGGS", 8.0)])
        self.assertFalse(any("Production Distribution" in title for title in titles))

    def test_complementary_chart_specs_use_filtered_rows_and_distinct_measures(self):
        inventory = _report_chart_specs({"report_type": "inventory", "rows": [
            {"species": "Cattle", "quantity": 4, "barangay": "Manggas"},
            {"species": "Goat", "quantity": 2, "barangay": "Manggas"},
            {"species": "Goat", "quantity": 3, "barangay": "Poblacion"},
        ]})
        composition = next(chart for chart in inventory if chart["title"] == "Species Composition by Barangay")
        self.assertEqual(composition["type"], "stacked_bar")
        self.assertEqual(composition["series"], ["Cattle", "Goat"])
        self.assertEqual(composition["data"][0], {"name": "Manggas", "values": {"Cattle": 4.0, "Goat": 2.0}})

        production = _report_chart_specs({"report_type": "production", "rows": [
            {"record_date": "2026-01-04", "production_type": "MILK", "quantity": 10, "unit": "LITERS", "barangay": "Manggas"},
            {"record_date": "2026-01-12", "production_type": "HONEY", "quantity": 3, "unit": "LITERS", "barangay": "Poblacion"},
            {"record_date": "2026-02-02", "production_type": "MILK", "quantity": 12, "unit": "LITERS", "barangay": "Manggas"},
            {"record_date": "2026-02-05", "production_type": "EGGS", "quantity": 8, "unit": "PIECES", "barangay": "Manggas"},
        ]})
        type_trend = next(chart for chart in production if chart["type"] == "multi_line")
        self.assertEqual(type_trend["series"], ["HONEY", "MILK"])
        self.assertEqual([month for month, _ in type_trend["data"]], ["2026-01", "2026-02"])
        self.assertFalse(any(chart.get("type") == "multi_line" and "EGGS" in chart.get("series", []) for chart in production))
        # Render a type/month gap too: the PDF must omit the missing value instead of plotting an invented zero.
        from reportlab.lib import colors as pdf_colors
        from reportlab.pdfgen.canvas import Canvas
        chart_pdf = BytesIO()
        pdf_canvas = Canvas(chart_pdf)
        flowable = _report_chart_flowable(type_trend, 640, 150, pdf_colors)
        flowable.wrap(640, 150)
        flowable.drawOn(pdf_canvas, 0, 0)
        pdf_canvas.save()
        self.assertTrue(chart_pdf.getvalue().startswith(b"%PDF"))
        records = next(chart for chart in production if chart["title"] == "Number of Production Records by Barangay")
        self.assertEqual(records["data"], [("Manggas", 3), ("Poblacion", 1)])

        disease = _report_chart_specs({"report_type": "disease_mortality", "rows": [
            {"record_kind": "DISEASE", "barangay": "Manggas", "affected_or_dead": 9},
            {"record_kind": "DISEASE", "barangay": "Manggas", "affected_or_dead": 2},
            {"record_kind": "DISEASE", "barangay": "Poblacion", "affected_or_dead": 1},
        ]})
        case_counts = next(chart for chart in disease if chart["title"] == "Reported Disease Cases by Barangay")
        self.assertEqual(case_counts["data"], [("Manggas", 2), ("Poblacion", 1)])

        slaughter = _report_chart_specs({"report_type": "slaughter", "rows": [
            {"species": "Cattle", "quantity": 2, "carcass_weight_kg": 360, "barangay": "Manggas"},
            {"species": "Cattle", "quantity": 1, "carcass_weight_kg": 150, "barangay": "Manggas"},
            {"species": "Goat", "quantity": 3, "carcass_weight_kg": None, "barangay": "Poblacion"},
        ]})
        average = next(chart for chart in slaughter if chart["title"] == "Average Recorded Carcass Weight by Species")
        self.assertEqual(average["data"], [("Cattle", 170.0)])
        by_barangay = next(chart for chart in slaughter if chart["title"] == "Slaughter by Barangay")
        self.assertEqual(by_barangay["data"], [("Manggas", 3), ("Poblacion", 3)])

        movement = _report_chart_specs({"report_type": "movement", "rows": [
            {"record_kind": "AUCTION", "record_date": "2026-01-05", "quantity": 5},
            {"record_kind": "MOVEMENT", "record_date": "2026-01-06", "control_number": "CLR-1", "quantity": 7},
            {"record_kind": "MOVEMENT", "record_date": "2026-01-07", "control_number": "CLR-1", "quantity": 2},
        ]})
        auction_records = next(chart for chart in movement if chart["title"] == "Auction Records Over Time")
        movement_records = next(chart for chart in movement if chart["title"] == "Movement Records Over Time")
        self.assertEqual(auction_records["data"], [("2026-01", 1)])
        self.assertEqual(movement_records["data"], [("2026-01", 1)])

    def test_pie_chart_spec_uses_actual_report_rows(self):
        inventory_report = {"report_type": "inventory", "rows": [
            {"species": "Cattle", "quantity": 4, "barangay": "Poblacion"},
            {"species": "Goat", "quantity": 2, "barangay": "Manggas"},
        ]}
        species_pie = next(chart for chart in _report_chart_specs(inventory_report)
                           if chart["title"] == "Livestock Distribution by Species")
        self.assertEqual(species_pie["type"], "pie")
        self.assertEqual(species_pie["data"], [("Cattle", 4), ("Goat", 2)])
        inventory_analysis = self.preview("inventory").data["analysis"]
        self.assertIn("Top 5 Barangays by Registered Livestock", [item["title"] for item in inventory_analysis["rankings"]])

    def test_pdf_chart_specs_show_empty_states_without_inventing_values(self):
        empty_report = self.client.get(
            reverse("reports-preview"),
            self.params("disease_mortality", date_from="2020-01-01", date_to="2020-01-31"),
        ).data
        charts = _report_chart_specs(empty_report)
        self.assertTrue(charts)
        self.assertTrue(all(chart["data"] == [] for chart in charts))

    def test_analysis_empty_coverage_and_unavailable_comparison_are_explicit(self):
        response = self.client.get(
            reverse("reports-preview"),
            self.params("disease_mortality", date_from="2020-01-01", date_to="2020-01-31"),
        )
        analysis = response.data["analysis"]
        self.assertEqual(analysis["coverage_level"], "none")
        self.assertIn("No approved records", analysis["coverage_notice"])
        self.assertFalse(analysis["comparison"]["available"])
        self.assertEqual(analysis["key_findings"], [])
        self.assertFalse(self.preview("production").data["analysis"]["comparison"]["available"])

    def test_previous_period_uses_same_duration_and_supports_comparison(self):
        previous_day = self.today - timedelta(days=1)
        ProductionRecord.objects.create(
            livestock=self.animal, production_type="MILK", quantity=Decimal("5"), unit="LITERS",
            record_date=previous_day, status="APPROVED", created_by=self.farmer_user,
        )
        report = self.preview("production", barangay="Poblacion").data
        comparison = report["analysis"]["comparison"]
        self.assertTrue(comparison["available"])
        self.assertEqual(comparison["period"], {"date_from": previous_day.isoformat(), "date_to": previous_day.isoformat()})
        milk = next(metric for metric in comparison["metrics"] if metric["label"] == "Liters")
        self.assertEqual((milk["current"], milk["previous"], milk["change_percent"]), (12.5, 5, 150.0))
        self.assertFalse(any(metric["label"] == "PIECES" for metric in comparison["metrics"]))

    def test_rankings_are_limited_to_five_and_sorted(self):
        rows = [{"species": "Cattle", "quantity": amount, "barangay": f"Barangay {amount}"}
                for amount in (1, 6, 3, 5, 2, 4)]
        analysis = _build_analysis("inventory", rows, [], {"date_from": "2025-01-01", "date_to": "2025-01-31"},
                                   {"records": 6, "total_heads": 21})
        ranking = analysis["rankings"][0]["items"]
        self.assertEqual(len(ranking), 5)
        self.assertEqual([item["value"] for item in ranking], [6, 5, 4, 3, 2])

    def test_pdf_month_axis_keeps_all_points_and_adapts_visible_ticks(self):
        short_months = ["2026-01", "2026-02", "2026-03", "2026-04"]
        self.assertEqual(list(_monthly_axis_ticks(short_months).values()), ["Jan 26", "Feb 26", "Mar 26", "Apr 26"])

        long_months = [f"{year:04d}-{month:02d}" for year in range(2022, 2026) for month in range(1, 13)]
        visible_ticks = _monthly_axis_ticks(long_months)
        self.assertLessEqual(len(visible_ticks), 8)
        self.assertIn(0, visible_ticks)
        self.assertIn(len(long_months) - 1, visible_ticks)
        # Tick selection only changes printed labels; the chart specification
        # still carries every month and therefore every line-chart point.
        report = {"report_type": "production", "rows": [
            {"record_date": f"{month}-15", "production_type": "MILK", "quantity": index + 1, "unit": "LITERS"}
            for index, month in enumerate(long_months)
        ]}
        trend = next(chart for chart in _report_chart_specs(report) if chart["type"] == "line")
        self.assertEqual(len(trend["data"]), len(long_months))

    def test_movement_direction_labels(self):
        from analytics.services.movement_direction import classify_movement_direction
        local = ["Poblacion", "Manggas"]
        self.assertEqual(classify_movement_direction("Lipa", "Poblacion, Padre Garcia", local), "INBOUND")
        self.assertEqual(classify_movement_direction("Manggas", "Lipa", local), "OUTBOUND")
        self.assertEqual(classify_movement_direction("Manggas", "Poblacion", local), "INTERNAL")
        self.assertEqual(classify_movement_direction("Tanauan", "Lipa", local), "UNKNOWN")
