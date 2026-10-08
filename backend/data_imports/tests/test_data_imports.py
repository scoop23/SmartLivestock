import io
import csv
from openpyxl import load_workbook
from datetime import date, timedelta
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APITestCase
from rest_framework import status

from livestock.models import Barangay, Farmer, LivestockType, LivestockInventory
from production.models import ProductionRecord, SlaughterRecord, LiveAnimalSale
from diseases.models import DiseaseCase, MortalityRecord
from users.models import User, Role
from data_imports.models import DataImportBatch
from data_imports.services.templates import generate_template_file
from data_imports.services.parser import parse_csv


class DataImportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        # 1. Barangays
        cls.manggas = Barangay.objects.create(
            barangay_name="Manggas", latitude=13.8821, longitude=121.2144
        )
        cls.banaba = Barangay.objects.create(
            barangay_name="Banaba", latitude=13.8750, longitude=121.2200
        )

        # 2. Roles
        cls.roles = {
            name: Role.objects.create(role_name=name)
            for name in ("FARMER", "SIBAT", "MAO", "ADMIN")
        }

        # 3. Users
        def create_user(username, role_name, barangay=None):
            return User.objects.create_user(
                username=username,
                email=f"{username}@smartlivestock.test",
                password="password123",
                role=cls.roles[role_name],
                account_status="APPROVED",
                assigned_barangay=barangay,
            )

        cls.mao_user = create_user("mao_officer", "MAO")
        cls.farmer_user = create_user("juan_farmer", "FARMER")
        cls.sibat_user = create_user("sibat_reviewer", "SIBAT", cls.manggas)

        # 4. Farmer profile
        cls.farmer = Farmer.objects.create(
            user=cls.farmer_user,
            barangay=cls.manggas,
            address="Purok 3, Manggas",
        )

        # 5. Livestock types
        cls.cattle = LivestockType.objects.create(name="Cattle")
        cls.swine = LivestockType.objects.create(name="Swine")

        # 6. Existing animal for duplicate check
        cls.existing_animal = LivestockInventory.objects.create(
            farmer=cls.farmer,
            livestock_type=cls.cattle,
            tag_number="EXISTING-CAT-001",
            entry_type="INDIVIDUAL",
            quantity=1,
            operational_status="ACTIVE",
            status="APPROVED",
            created_by=cls.mao_user,
        )

    def test_permission_denied_for_farmer_and_sibat(self):
        """Farmers and SIBAT inspectors must NOT be allowed to access bulk upload endpoints."""
        for unauthorized_user in [self.farmer_user, self.sibat_user]:
            self.client.force_authenticate(unauthorized_user)
            res = self.client.get("/api/data-imports/datasets/")
            self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
            res_val = self.client.post("/api/data-imports/validate/", {})
            self.assertEqual(res_val.status_code, status.HTTP_403_FORBIDDEN)

    def test_csv_parser_preserves_values_after_blank_header_columns(self):
        content = b"barangay,,livestock_type\nManggas,,Cattle\n"
        file = SimpleUploadedFile("blank_header.csv", content, content_type="text/csv")

        headers, rows, error = parse_csv(file)

        self.assertIsNone(error)
        self.assertEqual(headers, ["barangay", "livestock_type"])
        self.assertEqual(rows[0]["barangay"], "Manggas")
        self.assertEqual(rows[0]["livestock_type"], "Cattle")

    def test_invalid_inventory_and_auction_values_are_validation_errors(self):
        self.client.force_authenticate(self.mao_user)
        cases = [
            (
                "livestock_inventory",
                "Farmer,Barangay,Livestock Type,Entry Type,Sex\n"
                "juan_farmer,Manggas,Cattle,UNKNOWN,UNKNOWN\n",
                {"entry_type", "sex"},
            ),
            (
                "auction",
                "Farmer,Barangay,Quantity,Total Price,Sale Date,Sale Method,Purpose,Price Per Head\n"
                "juan_farmer,Manggas,1,1000,2026-04-18,INVALID,INVALID,not-a-number\n",
                {"sale_method", "purpose", "price_per_head"},
            ),
        ]

        for dataset_type, content, expected_fields in cases:
            with self.subTest(dataset_type=dataset_type):
                file = SimpleUploadedFile("invalid_values.csv", content.encode(), content_type="text/csv")
                response = self.client.post(
                    "/api/data-imports/validate/",
                    {"dataset_type": dataset_type, "file": file},
                    format="multipart",
                )
                self.assertEqual(response.status_code, status.HTTP_200_OK)
                fields = {issue["field"] for issue in response.data["preview_rows"][0]["issues"]}
                self.assertTrue(expected_fields.issubset(fields), fields)

    def test_list_datasets_for_mao(self):
        """MAO user can view all available import dataset schemas."""
        self.client.force_authenticate(self.mao_user)
        res = self.client.get("/api/data-imports/datasets/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        codes = [d["code"] for d in res.data]
        self.assertIn("livestock_inventory", codes)
        self.assertIn("production", codes)
        self.assertIn("disease", codes)
        self.assertIn("mortality", codes)
        self.assertIn("slaughter", codes)
        self.assertIn("auction", codes)
        slaughter = next(dataset for dataset in res.data if dataset["code"] == "slaughter")
        self.assertEqual(
            slaughter["required_fields"],
            ["barangay", "livestock_type", "quantity", "record_date"],
        )
        self.assertEqual(
            slaughter["description"],
            "Municipal abattoir and authorized on-farm slaughter events, including slaughter quantities and carcass yields.",
        )

    def test_validate_csv_with_valid_and_invalid_rows(self):
        """Validation returns counts of valid, warning, and error rows without mutating database."""
        self.client.force_authenticate(self.mao_user)

        csv_content = (
            "Farmer,Barangay,Livestock_Type,Tag Number,Entry Type,Quantity,Breed,Sex,Weight\n"
            # Row 2: Valid
            "juan_farmer,Manggas,Cattle,NEW-CAT-002,INDIVIDUAL,1,Brahman,MALE,450.00\n"
            # Row 3: Missing required barangay
            "juan_farmer,,Cattle,NEW-CAT-003,INDIVIDUAL,1,Brahman,MALE,420.00\n"
            # Row 4: Invalid foreign key (Unknown Barangay)
            "juan_farmer,NonExistentBarangay,Cattle,NEW-CAT-004,INDIVIDUAL,1,Brahman,MALE,400.00\n"
            # Row 5: Duplicate warning (existing tag in DB)
            "juan_farmer,Manggas,Cattle,EXISTING-CAT-001,INDIVIDUAL,1,Brahman,MALE,410.00\n"
            # Row 6: Invalid constraint (INDIVIDUAL quantity cannot be > 1)
            "juan_farmer,Manggas,Cattle,NEW-CAT-006,INDIVIDUAL,5,Brahman,MALE,300.00\n"
        )
        file = SimpleUploadedFile("inventory_test.csv", csv_content.encode("utf-8"), content_type="text/csv")

        res = self.client.post(
            "/api/data-imports/validate/",
            {"dataset_type": "livestock_inventory", "file": file},
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["total_rows"], 5)
        self.assertEqual(res.data["valid_count"], 1)
        self.assertEqual(res.data["warning_count"], 1)  # EXISTING-CAT-001
        self.assertEqual(res.data["error_count"], 3)  # Missing, Unknown brgy, Qty>1 for individual
        self.assertEqual(res.data["column_mapping"]["livestock_type"], "Livestock_Type")
        self.assertEqual(res.data["preview_rows"][0]["data"]["livestock_type"], "Cattle")

        # Verify database was NOT touched
        self.assertFalse(LivestockInventory.objects.filter(tag_number="NEW-CAT-002").exists())

    def test_execute_batch_import_transaction_safety(self):
        """Valid rows are imported into database and assigned official MAO approval metadata."""
        self.client.force_authenticate(self.mao_user)

        csv_content = (
            "Farmer,Barangay,Livestock_Type,Tag Number,Entry Type,Quantity,Breed,Sex,Weight\n"
            "juan_farmer,Manggas,Cattle,IMPORT-CAT-101,INDIVIDUAL,1,Brahman,MALE,460.00\n"
            "juan_farmer,Manggas,Cattle,IMPORT-CAT-102,INDIVIDUAL,1,Brahman,FEMALE,420.00\n"
            # Duplicate tag warning
            "juan_farmer,Manggas,Cattle,EXISTING-CAT-001,INDIVIDUAL,1,Brahman,MALE,410.00\n"
        )
        file = SimpleUploadedFile("import_exec.csv", csv_content.encode("utf-8"), content_type="text/csv")

        res = self.client.post(
            "/api/data-imports/import/",
            {
                "dataset_type": "livestock_inventory",
                "file": file,
                "skip_duplicates": True,
                "target_status": "APPROVED",
            },
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["imported_rows"], 2)
        self.assertEqual(res.data["skipped_rows"], 1)

        # Check DB persistence
        imported_animal = LivestockInventory.objects.get(tag_number="IMPORT-CAT-101")
        self.assertEqual(imported_animal.livestock_type, self.cattle)
        self.assertEqual(imported_animal.status, "APPROVED")
        self.assertEqual(imported_animal.reviewed_by, self.mao_user)
        self.assertIsNotNone(imported_animal.reviewed_at)
        self.assertIn("Historical bulk import by MAO", imported_animal.review_remarks)
        self.assertEqual(
            LivestockInventory.objects.get(tag_number="IMPORT-CAT-102").sex,
            "FEMALE",
        )

        # Check batch record
        batch = DataImportBatch.objects.get(pk=res.data["batch_id"])
        self.assertEqual(batch.imported_rows, 2)
        self.assertEqual(batch.skipped_rows, 1)
        self.assertEqual(batch.status, DataImportBatch.ImportStatus.PARTIAL)

    def test_legacy_species_inventory_header_still_imports_to_livestock_type(self):
        """Older inventory sheets using Species remain aliases for the same FK."""
        self.client.force_authenticate(self.mao_user)
        csv_content = (
            "Farmer,Barangay,Species,Tag Number,Entry Type,Quantity\n"
            "juan_farmer,Manggas,Cattle,LEGACY-CAT-101,INDIVIDUAL,1\n"
        )
        file = SimpleUploadedFile("legacy_inventory.csv", csv_content.encode("utf-8"), content_type="text/csv")

        response = self.client.post(
            "/api/data-imports/import/",
            {"dataset_type": "livestock_inventory", "file": file, "target_status": "APPROVED"},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported_rows"], 1)
        animal = LivestockInventory.objects.get(tag_number="LEGACY-CAT-101")
        self.assertEqual(animal.livestock_type, self.cattle)

    def test_slaughter_import_accepts_canonical_and_legacy_type_headers(self):
        self.client.force_authenticate(self.mao_user)
        for header, tag in (("livestock_type", "TYPE"), ("species", "LEGACY")):
            csv_content = (
                f"barangay,{header},quantity,record_date,carcass_weight\n"
                f"Manggas,Cattle,2,2026-04-15,420.50\n"
            )
            file = SimpleUploadedFile(
                f"slaughter_{tag.lower()}.csv",
                csv_content.encode("utf-8"),
                content_type="text/csv",
            )
            response = self.client.post(
                "/api/data-imports/import/",
                {"dataset_type": "slaughter", "file": file, "target_status": "APPROVED"},
                format="multipart",
            )

            self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
            self.assertEqual(response.data["imported_rows"], 1)

        self.assertEqual(SlaughterRecord.objects.count(), 2)
        self.assertEqual(
            set(SlaughterRecord.objects.values_list("livestock_type_id", flat=True)),
            {self.cattle.id},
        )

    def test_slaughter_validation_uses_livestock_type_language(self):
        self.client.force_authenticate(self.mao_user)
        csv_content = (
            "barangay,livestock_type,quantity,record_date\n"
            "Manggas,Unknown Animal,1,2026-04-15\n"
        )
        file = SimpleUploadedFile("slaughter_invalid.csv", csv_content.encode("utf-8"), content_type="text/csv")

        response = self.client.post(
            "/api/data-imports/validate/",
            {"dataset_type": "slaughter", "file": file},
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["column_mapping"]["livestock_type"], "livestock_type")
        issue = response.data["preview_rows"][0]["issues"][0]
        self.assertEqual(issue["field"], "livestock_type")
        self.assertIn("Livestock Type", issue["message"])
        self.assertEqual(response.data["issues_sample"][0]["livestock_type"], "Unknown Animal")
        self.assertNotIn("species", response.data["issues_sample"][0])

    def test_slaughter_templates_use_livestock_type_headers(self):
        self.client.force_authenticate(self.mao_user)

        csv_response = self.client.get("/api/data-imports/templates/slaughter/?format=csv")
        self.assertEqual(csv_response.status_code, status.HTTP_200_OK)
        csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
        self.assertEqual(csv_rows[0][:4], ["barangay", "livestock_type", "quantity", "carcass_weight"])
        self.assertNotIn("species", csv_rows[0])

        xlsx_response = self.client.get("/api/data-imports/templates/slaughter/?format=xlsx")
        self.assertEqual(xlsx_response.status_code, status.HTTP_200_OK)
        workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
        headers = [cell.value for cell in workbook.active[1]]
        self.assertIn("livestock_type", headers)
        self.assertNotIn("species", headers)
        workbook.close()

    def test_slaughter_error_report_uses_livestock_type_header_and_reads_legacy_logs(self):
        self.client.force_authenticate(self.mao_user)
        batch = DataImportBatch.objects.create(
            dataset_type="slaughter",
            file_name="old_slaughter.csv",
            uploaded_by=self.mao_user,
            total_rows=1,
            error_rows=1,
            error_log=[{
                "row_number": 2,
                "barangay": "Manggas",
                "species": "Unknown Animal",
                "field": "species",
                "error_type": "INVALID_FOREIGN_KEY",
                "severity": "ERROR",
                "error_message": "Species 'Unknown Animal' is not recognized. Valid species: Cattle.",
            }],
        )

        response = self.client.get(f"/api/data-imports/batches/{batch.id}/errors/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        content = response.content.decode("utf-8-sig")
        self.assertIn("Row Number,Barangay,Livestock Type,Field,Error Type,Severity,Error Message", content)
        self.assertIn(
            "2,Manggas,Unknown Animal,Livestock Type,INVALID_FOREIGN_KEY,ERROR,"
            "Livestock Type 'Unknown Animal' is not recognized. Valid livestock type: Cattle.",
            content,
        )

    def test_production_and_disease_and_mortality_imports(self):
        """Test importing Production, Disease, and Mortality records."""
        self.client.force_authenticate(self.mao_user)

        # 1. Production
        prod_csv = (
            "Farmer,Barangay,Commodity,Quantity,Unit,Date,Notes\n"
            "juan_farmer,Manggas,MILK,30.5,LITERS,2026-03-25,Morning milk\n"
        )
        file_prod = SimpleUploadedFile("prod.csv", prod_csv.encode("utf-8"), content_type="text/csv")
        res_prod = self.client.post(
            "/api/data-imports/import/",
            {"dataset_type": "production", "file": file_prod},
            format="multipart",
        )
        self.assertEqual(res_prod.status_code, status.HTTP_200_OK)
        self.assertEqual(res_prod.data["imported_rows"], 1)
        self.assertTrue(ProductionRecord.objects.filter(quantity=30.5, status="APPROVED").exists())

        # 2. Disease
        dis_csv = (
            "Farmer,Barangay,Disease Name,Affected Count,Date\n"
            "juan_farmer,Manggas,Foot and Mouth Disease,3,2026-03-26\n"
        )
        file_dis = SimpleUploadedFile("dis.csv", dis_csv.encode("utf-8"), content_type="text/csv")
        res_dis = self.client.post(
            "/api/data-imports/import/",
            {"dataset_type": "disease", "file": file_dis},
            format="multipart",
        )
        self.assertEqual(res_dis.status_code, status.HTTP_200_OK)
        self.assertEqual(res_dis.data["imported_rows"], 1)
        self.assertTrue(DiseaseCase.objects.filter(name="Foot and Mouth Disease", affected_count=3).exists())

        # 3. Mortality
        mort_csv = (
            "Farmer,Barangay,Cause,Death Count,Date\n"
            "juan_farmer,Manggas,Heat Stroke,1,2026-03-27\n"
        )
        file_mort = SimpleUploadedFile("mort.csv", mort_csv.encode("utf-8"), content_type="text/csv")
        res_mort = self.client.post(
            "/api/data-imports/import/",
            {"dataset_type": "mortality", "file": file_mort},
            format="multipart",
        )
        self.assertEqual(res_mort.status_code, status.HTTP_200_OK)
        self.assertEqual(res_mort.data["imported_rows"], 1)
        self.assertTrue(MortalityRecord.objects.filter(cause="Heat Stroke", death_count=1).exists())

    def test_download_template_xlsx_and_csv(self):
        """MAO can download official formatted templates in XLSX or CSV."""
        self.client.force_authenticate(self.mao_user)

        # CSV format
        res_csv = self.client.get("/api/data-imports/templates/livestock_inventory/?format=csv")
        self.assertEqual(res_csv.status_code, status.HTTP_200_OK)
        self.assertIn("text/csv", res_csv["Content-Type"])
        self.assertIn("Livestock_Inventory_Template.csv", res_csv["Content-Disposition"])
        csv_rows = list(csv.reader(io.StringIO(res_csv.content.decode("utf-8-sig"))))
        self.assertIn("livestock_type", csv_rows[0])
        self.assertNotIn("species", csv_rows[0])
        livestock_type_column = csv_rows[0].index("livestock_type")
        self.assertEqual(csv_rows[1][livestock_type_column], "Cattle")

        # XLSX format
        res_xlsx = self.client.get("/api/data-imports/templates/livestock_inventory/?format=xlsx")
        self.assertEqual(res_xlsx.status_code, status.HTTP_200_OK)
        self.assertIn("application/vnd.openxmlformats", res_xlsx["Content-Type"])
        workbook = load_workbook(io.BytesIO(res_xlsx.content), read_only=True)
        sheet = workbook.active
        xlsx_headers = [cell.value for cell in sheet[1]]
        self.assertIn("livestock_type", xlsx_headers)
        self.assertNotIn("species", xlsx_headers)
        self.assertEqual(sheet.cell(row=2, column=xlsx_headers.index("livestock_type") + 1).value, "Cattle")
        workbook.close()

    def test_download_error_report(self):
        """MAO can download a structured CSV error report for any batch."""
        self.client.force_authenticate(self.mao_user)
        batch = DataImportBatch.objects.create(
            dataset_type="livestock_inventory",
            file_name="bad_upload.csv",
            uploaded_by=self.mao_user,
            total_rows=2,
            error_rows=2,
            error_log=[
                {
                    "row_number": 3,
                    "barangay": "Unknown",
                    "species": "Cattle",  # Older stored logs remain readable.
                    "field": "barangay",
                    "error_type": "INVALID_FOREIGN_KEY",
                    "severity": "ERROR",
                    "error_message": "Barangay Unknown does not exist in Padre Garcia.",
                }
            ],
        )

        res = self.client.get(f"/api/data-imports/batches/{batch.id}/errors/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("text/csv", res["Content-Type"])
        content = res.content.decode("utf-8-sig")
        self.assertIn("Row Number,Barangay,Livestock Type,Field,Error Type,Severity,Error Message", content)
        self.assertIn("3,Unknown,Cattle,barangay,INVALID_FOREIGN_KEY,ERROR,Barangay Unknown does not exist in Padre Garcia.", content)

    def test_analytics_and_overview_reflection(self):
        """Imported APPROVED records must immediately reflect in Data Overview and Population analytics."""
        self.client.force_authenticate(self.mao_user)
        from analytics.services.overview import overview_summary

        # Check baseline overview
        initial_overview = overview_summary(user=self.mao_user)
        initial_heads = initial_overview["population"]["total_heads"]

        # Import 2 new heads as APPROVED
        csv_content = (
            "Farmer,Barangay,Species,Tag Number,Entry Type,Quantity,Breed,Sex,Weight\n"
            "juan_farmer,Manggas,Cattle,POP-CAT-901,INDIVIDUAL,1,Brahman,MALE,400.00\n"
            "juan_farmer,Manggas,Cattle,POP-CAT-902,INDIVIDUAL,1,Brahman,FEMALE,380.00\n"
        )
        file = SimpleUploadedFile("pop_test.csv", csv_content.encode("utf-8"), content_type="text/csv")
        res = self.client.post(
            "/api/data-imports/import/",
            {"dataset_type": "livestock_inventory", "file": file, "target_status": "APPROVED"},
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        # Re-check overview: heads must have increased by exactly 2
        updated_overview = overview_summary(user=self.mao_user)
        self.assertEqual(updated_overview["population"]["total_heads"], initial_heads + 2)

    def test_batch_processing_performance(self):
        """Verify bulk batch processing handles 100 generated rows efficiently."""
        self.client.force_authenticate(self.mao_user)
        lines = ["Farmer,Barangay,Commodity,Quantity,Unit,Date,Notes"]
        for i in range(1, 101):
            lines.append(f"juan_farmer,Manggas,MILK,15.5,LITERS,2026-03-20,Milk entry #{i}")
        csv_content = "\n".join(lines)

        file = SimpleUploadedFile("large_prod.csv", csv_content.encode("utf-8"), content_type="text/csv")
        res = self.client.post(
            "/api/data-imports/import/",
            {"dataset_type": "production", "file": file},
            format="multipart",
        )
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["imported_rows"], 100)
        self.assertLess(res.data["duration_seconds"], 5.0)  # Must complete in < 5 seconds
