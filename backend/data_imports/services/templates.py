import csv
import io
from typing import Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .datasets import get_dataset_config


def generate_template_file(dataset_type: str, file_format: str = "xlsx") -> Tuple[bytes, str, str]:
    """
    Generate an official downloadable municipal template (.xlsx or .csv)
    for a supported dataset.
    Returns:
        (file_bytes, filename, content_type)
    """
    config = get_dataset_config(dataset_type)
    headers = list(config.field_aliases.keys())
    sample_rows = config.sample_rows

    if file_format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in sample_rows:
            writer.writerow([row.get(h, "") for h in headers])
        
        content = output.getvalue().encode("utf-8-sig")
        filename = f"{config.label.replace(' ', '_')}_Template.csv"
        content_type = "text/csv; charset=utf-8"
        return content, filename, content_type

    # Excel (.xlsx) format with rich styling
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.label[:31]  # Excel 31-char sheet name limit

    # Professional styles
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E4D2B", end_color="1E4D2B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sample_font = Font(name="Arial", size=10, italic=True, color="475569")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Row 1: Headers
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Rows 2+: Sample data
    for row_idx, sample in enumerate(sample_rows, start=2):
        row_values = [sample.get(h, "") for h in headers]
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 22
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = sample_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    stream = io.BytesIO()
    wb.save(stream)
    content = stream.getvalue()
    wb.close()

    filename = f"{config.label.replace(' ', '_')}_Template.xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return content, filename, content_type
